"""Bounded local text extraction; deliberately does not claim OCR support."""

import asyncio
import json
import sys
import zipfile
from io import BytesIO
from pathlib import Path

import psutil

MAX_UPLOAD = 10 * 1024 * 1024
MAX_TEXT = 2_000_000
MAX_PAGES = 300
MAX_ZIP_SIZE = 50 * 1024 * 1024
MAX_MEMORY = 768 * 1024 * 1024


def extract(name: str, content: bytes) -> list[dict]:
    if len(content) > MAX_UPLOAD:
        raise ValueError("文件超过 10 MiB 限额。")
    extension = name.rsplit(".", 1)[-1].lower()
    pages: list[tuple[int | None, str]] = []
    if extension in ("txt", "md"):
        try:
            pages.append((None, content.decode("utf-8-sig")))
        except UnicodeDecodeError as exc:
            raise ValueError("文本文件必须使用 UTF-8 编码。") from exc
    elif extension == "pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise ValueError("暂不支持加密 PDF。")
            if len(reader.pages) > MAX_PAGES:
                raise ValueError("PDF 超过 300 页限额。")
            total = 0
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                total += len(text)
                if total > MAX_TEXT:
                    raise ValueError("提取文本超过 200 万字符限额。")
                pages.append((number, text))
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF 无法解析；请确认文件完整且包含可提取文字。") from exc
    elif extension == "docx":
        from docx import Document

        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > 2000
                    or sum(item.file_size for item in entries) > MAX_ZIP_SIZE
                ):
                    raise ValueError("DOCX 解压规模超过限制。")
                if any(
                    item.flag_bits & 1
                    or item.file_size > 10 * 1024 * 1024
                    or item.file_size / max(1, item.compress_size) > 200
                    for item in entries
                ):
                    raise ValueError("DOCX 存在加密或异常压缩内容。")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("文件不是有效的 DOCX 文档。")
            document = Document(BytesIO(content))
            segments = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                segments.extend(
                    " | ".join(cell.text for cell in row.cells) for row in table.rows
                )
            pages.append((None, "\n".join(segments)))
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("DOCX 无法解析；请确认文件完整。") from exc
    else:
        raise ValueError("只支持 PDF、DOCX、TXT 和 MD 文件。")
    if sum(len(text) for _, text in pages) > MAX_TEXT:
        raise ValueError("提取文本超过 200 万字符限额。")
    if not any(text.strip() for _, text in pages):
        raise ValueError("未提取到文字；扫描 PDF 需要另行 OCR，本入口未提供 OCR。")
    chunks = []
    for page, text in pages:
        text = text.replace("\x00", "").strip()
        for offset in range(0, len(text), 3000):
            chunks.append(
                {
                    "id": str(len(chunks) + 1),
                    "text": text[offset : offset + 3000],
                    "page": page,
                }
            )
    return chunks


async def parse_upload(name: str, content: bytes) -> list[dict]:
    """在独立进程解析，限制时间、CPU、文本规模并监控内存使用。"""
    process = await asyncio.create_subprocess_exec(
        sys.executable,
        str(Path(__file__).resolve()),
        name,
        stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.DEVNULL,
    )

    async def watch_memory():
        while process.returncode is None:
            try:
                if psutil.Process(process.pid).memory_info().rss > MAX_MEMORY:
                    raise ValueError("解析内存超过 768 MiB 限额；请拆分文件后重试。")
            except psutil.NoSuchProcess:
                return
            except psutil.Error as exc:
                raise ValueError("无法监控解析进程的资源使用，已停止解析。") from exc
            await asyncio.sleep(0.02)

    communication = asyncio.create_task(process.communicate(content))
    monitor = asyncio.create_task(watch_memory())
    try:
        async with asyncio.timeout(30):
            done, _ = await asyncio.wait(
                (communication, monitor), return_when=asyncio.FIRST_COMPLETED
            )
            if monitor in done:
                monitor.result()
            stdout, _ = await communication
    except (ValueError, TimeoutError, asyncio.CancelledError) as exc:
        if process.returncode is None:
            process.kill()
        await process.wait()
        communication.cancel()
        await asyncio.gather(communication, return_exceptions=True)
        if isinstance(exc, asyncio.CancelledError):
            raise
        if isinstance(exc, ValueError):
            raise
        raise ValueError("解析超过 30 秒限制；请拆分或简化文件后重试。") from exc
    finally:
        monitor.cancel()
        await asyncio.gather(monitor, return_exceptions=True)
    if process.returncode != 0:
        raise ValueError("解析进程失败或达到资源限制；请拆分文件后重试。")
    try:
        result = json.loads(stdout)
    except ValueError as exc:
        raise ValueError("解析进程返回无效结果。") from exc
    if "error" in result:
        raise ValueError(result["error"])
    return result["chunks"]


if __name__ == "__main__":
    import resource

    resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
    # macOS rejects lowering AS/DATA limits on this runtime; the parent monitors RSS.
    if sys.platform != "darwin":
        resource.setrlimit(resource.RLIMIT_AS, (MAX_MEMORY, MAX_MEMORY))
    try:
        result = {"chunks": extract(sys.argv[1], sys.stdin.buffer.read(MAX_UPLOAD + 1))}
    except ValueError as exc:
        result = {"error": str(exc)[:500]}
    except Exception:  # noqa: BLE001 -- Worker protocol must return only a safe diagnostic for arbitrary parser failures.
        result = {"error": "文档无法解析；请确认文件格式与完整性。"}
    sys.stdout.write(json.dumps(result, ensure_ascii=False))
