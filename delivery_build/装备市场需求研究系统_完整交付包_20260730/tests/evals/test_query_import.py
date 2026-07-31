from pathlib import Path
import json

from openpyxl import Workbook

from evals.cli import validate_query_file
from evals.models import read_jsonl
from evals.query_import import import_expert_workbook, load_eligible_candidates


HEADERS = [
    "候选ID", "候选Query", "态势核心", "区域", "领域", "威胁/压力", "难度建议", "模式", "划分",
    "场景说明", "主来源ID", "主来源链接", "辅来源ID", "辅来源链接", "来源触发依据", "专家A状态",
    "专家A修改后Query", "专家A理由/淘汰原因", "专家B复核", "是否仲裁", "最终Query", "仲裁/最终说明",
    "Query质量评分", "可研究性",
]


def _workbook(path: Path, count: int = 140) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "专家标注"
    sheet.append(["title"])
    sheet.append([])
    sheet.append([])
    sheet.append(HEADERS)
    for index in range(1, count + 1):
        row = {header: "" for header in HEADERS}
        row.update(
            {
                "候选ID": f"JST-A-{index:03d}",
                "候选Query": f"专家候选研究问题 {index}",
                "区域": ["台海", "南海", "西太"][index % 3],
                "领域": ["电磁", "低空", "海域", "无人"][index % 4],
                "难度建议": ["easy", "medium", "hard"][index % 3],
                "主来源链接": f"https://source.example/{index}",
                "辅来源链接": f"https://second.example/{index}",
                "专家A状态": "入选",
                "专家B复核": "同意",
                "是否仲裁": "否",
                "Query质量评分": "4-较好",
                "可研究性": "可研究",
            }
        )
        sheet.append([row[header] for header in HEADERS])
    workbook.save(path)
    return path


def test_import_expert_workbook_emits_minimal_public_rows(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "queries.xlsx")
    manifest = import_expert_workbook(source, tmp_path / "dataset")
    assert manifest["query_count"] == 132
    rows = read_jsonl(tmp_path / "dataset" / "queries.jsonl")
    assert set(rows[0]) == {"query_id", "query", "region", "domain", "difficulty", "split"}
    assert rows[0]["query_id"] == "Q-0001"
    assert sum(row["split"] == "pilot" for row in rows) == 12
    assert sum(row["split"] == "test" for row in rows) == 120
    assert "source" not in json.dumps(rows, ensure_ascii=False).lower()
    admin = read_jsonl(tmp_path / "dataset" / "admin_mapping.jsonl")
    assert admin[0]["source_candidate_id"].startswith("JST-")
    assert admin[0]["primary_source_url"].startswith("https://")
    assert validate_query_file(tmp_path / "dataset" / "queries.jsonl")["query_count"] == 132


def test_score_three_requires_final_query(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "queries.xlsx", count=2)
    workbook = __import__("openpyxl").load_workbook(source)
    sheet = workbook["专家标注"]
    score_col = HEADERS.index("Query质量评分") + 1
    final_col = HEADERS.index("最终Query") + 1
    sheet.cell(5, score_col, "3-可用需改")
    sheet.cell(6, score_col, "3-可用需改")
    sheet.cell(6, final_col, "专家最终改写问题")
    workbook.save(source)
    eligible = load_eligible_candidates(source)
    assert [item.query for item in eligible] == ["专家最终改写问题"]


def test_unreviewed_rows_are_not_eligible(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "queries.xlsx", count=1)
    workbook = __import__("openpyxl").load_workbook(source)
    sheet = workbook["专家标注"]
    sheet.cell(5, HEADERS.index("专家B复核") + 1, "未复核")
    workbook.save(source)
    assert load_eligible_candidates(source) == []


def test_unreviewed_rows_can_form_local_candidate_test_set(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "queries.xlsx", count=4)
    workbook = __import__("openpyxl").load_workbook(source)
    sheet = workbook["专家标注"]
    for row_number in range(5, 9):
        for header in (
            "专家A状态",
            "专家B复核",
            "是否仲裁",
            "Query质量评分",
            "可研究性",
        ):
            sheet.cell(row_number, HEADERS.index(header) + 1, "")
    workbook.save(source)
    output = tmp_path / "candidate-dataset"
    manifest = import_expert_workbook(
        source,
        output,
        total=4,
        pilot_count=1,
        require_expert_approval=False,
    )
    assert manifest["selection_mode"] == "local_candidate_test"
    assert manifest["query_count"] == 4
    assert [row["query_id"] for row in read_jsonl(output / "queries.jsonl")] == [
        "Q-0001",
        "Q-0002",
        "Q-0003",
        "Q-0004",
    ]
