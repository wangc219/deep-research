from __future__ import annotations

from collections import defaultdict, deque
from dataclasses import dataclass
from difflib import SequenceMatcher
from hashlib import sha256
from pathlib import Path
from typing import Any, Iterable
import json
import random
import re

from .models import EvalQuery, write_jsonl


SHEET_NAME = "专家标注"
HEADER_ROW = 4
ALLOWED_DIFFICULTIES = {"easy", "medium", "hard"}


@dataclass(frozen=True)
class ExpertCandidate:
    candidate_id: str
    query: str
    region: str
    domain: str
    difficulty: str
    admin: dict[str, Any]


def import_expert_workbook(
    workbook_path: str | Path,
    output_dir: str | Path,
    *,
    total: int = 132,
    pilot_count: int = 12,
    seed: int = 20260718,
    require_expert_approval: bool = True,
) -> dict[str, Any]:
    candidates = load_eligible_candidates(
        workbook_path,
        require_expert_approval=require_expert_approval,
    )
    selected = stratified_select(candidates, total=total, pilot_count=pilot_count, seed=seed)
    queries: list[EvalQuery] = []
    admin_rows: list[dict[str, Any]] = []
    for index, candidate in enumerate(selected, start=1):
        query_id = f"Q-{index:04d}"
        split = "pilot" if index <= pilot_count else "test"
        query = EvalQuery(
            query_id=query_id,
            query=candidate.query,
            region=candidate.region,
            domain=candidate.domain,
            difficulty=candidate.difficulty,
            split=split,
        )
        query.validate()
        queries.append(query)
        admin_rows.append({"query_id": query_id, **candidate.admin})

    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    public_path = output / "queries.jsonl"
    admin_path = output / "admin_mapping.jsonl"
    write_jsonl(public_path, [item.to_dict() for item in queries])
    write_jsonl(admin_path, admin_rows)
    near_duplicates = find_near_duplicates(selected)
    (output / "near_duplicates.json").write_text(
        json.dumps(
            [
                {"left_candidate_id": left, "right_candidate_id": right, "similarity": score}
                for left, right, score in near_duplicates
            ],
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    manifest = {
        "schema_version": "1.0",
        "source_workbook": str(Path(workbook_path).resolve()),
        "query_count": len(queries),
        "pilot_count": sum(item.split == "pilot" for item in queries),
        "test_count": sum(item.split == "test" for item in queries),
        "seed": seed,
        "public_sha256": sha256(public_path.read_bytes()).hexdigest(),
        "admin_sha256": sha256(admin_path.read_bytes()).hexdigest(),
        "near_duplicate_count": len(near_duplicates),
        "selection_mode": (
            "expert_approved" if require_expert_approval else "local_candidate_test"
        ),
        "fields": ["query_id", "query", "region", "domain", "difficulty", "split"],
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    return manifest


def load_eligible_candidates(
    workbook_path: str | Path,
    *,
    require_expert_approval: bool = True,
) -> list[ExpertCandidate]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment-specific guard
        raise RuntimeError("Excel import requires the optional eval dependency openpyxl") from exc

    workbook = load_workbook(Path(workbook_path), read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"workbook is missing sheet: {SHEET_NAME}")
    sheet = workbook[SHEET_NAME]
    headers = [str(cell.value or "").strip() for cell in sheet[HEADER_ROW]]
    index = {name: position for position, name in enumerate(headers) if name}
    required = {
        "候选ID", "候选Query", "区域", "领域", "难度建议", "专家A状态",
        "专家A修改后Query", "专家B复核", "是否仲裁", "最终Query",
        "Query质量评分", "可研究性",
    }
    missing = sorted(required - set(index))
    if missing:
        raise ValueError(f"expert workbook is missing columns: {', '.join(missing)}")

    result: list[ExpertCandidate] = []
    seen_queries: set[str] = set()
    for row in sheet.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        values = {name: _clean(row[position] if position < len(row) else "") for name, position in index.items()}
        if not values.get("候选ID"):
            continue
        if require_expert_approval and not _is_approved(values):
            continue
        query = _select_query_text(values)
        normalized = normalize_query(query)
        if not normalized or normalized in seen_queries:
            continue
        difficulty = str(values.get("难度建议", "medium")).lower()
        if difficulty not in ALLOWED_DIFFICULTIES:
            difficulty = "medium"
        seen_queries.add(normalized)
        result.append(
            ExpertCandidate(
                candidate_id=str(values["候选ID"]),
                query=query,
                region=str(values.get("区域") or "未标注"),
                domain=str(values.get("领域") or "未标注"),
                difficulty=difficulty,
                admin={
                    "source_candidate_id": str(values["候选ID"]),
                    "candidate_query": str(values.get("候选Query", "")),
                    "selected_query": query,
                    "expert_a_status": str(values.get("专家A状态", "")),
                    "expert_a_reason": str(values.get("专家A理由/淘汰原因", "")),
                    "expert_b_review": str(values.get("专家B复核", "")),
                    "arbitrated": str(values.get("是否仲裁", "")),
                    "final_note": str(values.get("仲裁/最终说明", "")),
                    "quality_score": str(values.get("Query质量评分", "")),
                    "researchability": str(values.get("可研究性", "")),
                    "primary_source_url": str(values.get("主来源链接", "")),
                    "secondary_source_url": str(values.get("辅来源链接", "")),
                },
            )
        )
    return result


def stratified_select(
    candidates: list[ExpertCandidate],
    *,
    total: int,
    pilot_count: int,
    seed: int,
) -> list[ExpertCandidate]:
    if pilot_count < 0 or total <= 0 or pilot_count > total:
        raise ValueError("invalid total/pilot counts")
    if len(candidates) < total:
        raise ValueError(f"only {len(candidates)} eligible expert queries; {total} required")
    rng = random.Random(seed)
    strata: dict[tuple[str, str, str], list[ExpertCandidate]] = defaultdict(list)
    for item in candidates:
        strata[(item.region, item.domain, item.difficulty)].append(item)
    queues: list[deque[ExpertCandidate]] = []
    for key in sorted(strata):
        values = list(strata[key])
        rng.shuffle(values)
        queues.append(deque(values))
    rng.shuffle(queues)

    selected: list[ExpertCandidate] = []
    while queues and len(selected) < total:
        remaining: list[deque[ExpertCandidate]] = []
        for queue in queues:
            if queue and len(selected) < total:
                selected.append(queue.popleft())
            if queue:
                remaining.append(queue)
        queues = remaining
    if len(selected) != total:
        raise RuntimeError("stratified selection did not produce the requested count")
    return selected


def find_near_duplicates(candidates: Iterable[ExpertCandidate], *, threshold: float = 0.92) -> list[tuple[str, str, float]]:
    values = list(candidates)
    duplicates: list[tuple[str, str, float]] = []
    for left_index, left in enumerate(values):
        left_text = normalize_query(left.query)
        for right in values[left_index + 1 :]:
            score = SequenceMatcher(None, left_text, normalize_query(right.query)).ratio()
            if score >= threshold:
                duplicates.append((left.candidate_id, right.candidate_id, round(score, 4)))
    return duplicates


def normalize_query(value: str) -> str:
    return re.sub(r"[^\w\u4e00-\u9fff]+", "", str(value).lower())


def _is_approved(values: dict[str, str]) -> bool:
    if values.get("专家A状态") == "淘汰":
        return False
    if values.get("可研究性") != "可研究":
        return False
    score = _parse_score(values.get("Query质量评分", ""))
    if score < 3:
        return False
    review = values.get("专家B复核")
    arbitration_passed = values.get("是否仲裁") == "是" and bool(values.get("最终Query"))
    if review != "同意" and not arbitration_passed:
        return False
    if score == 3 and not values.get("最终Query"):
        return False
    return bool(_select_query_text(values))


def _select_query_text(values: dict[str, str]) -> str:
    return str(values.get("最终Query") or values.get("专家A修改后Query") or values.get("候选Query") or "").strip()


def _parse_score(value: str) -> int:
    match = re.match(r"\s*([1-5])", str(value))
    return int(match.group(1)) if match else 0


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()
