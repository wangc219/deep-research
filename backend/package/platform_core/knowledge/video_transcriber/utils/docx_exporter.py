"""Export a video's original description and cleaned transcript as Word text."""

from pathlib import Path
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from platform_core.knowledge.video_transcriber.utils.transcript_layout import transcript_presentation


def video_docx_bytes(description: str, transcript: str, account: str = "") -> bytes:
    """生成带自然段、标题层级和稳定分页的可编辑 Word。"""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)

    for name, size in (("Normal", 11), ("Title", 20), ("Heading 1", 14), ("Subtitle", 10)):
        style = document.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attribute}"), None)
        fonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        fonts.set(qn("w:cs"), "Noto Sans CJK SC")
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    normal = document.styles["Normal"].paragraph_format
    normal.line_spacing = Pt(22)
    normal.space_after = Pt(10)
    normal.widow_control = True
    heading = document.styles["Heading 1"].paragraph_format
    heading.space_before, heading.space_after = Pt(18), Pt(10)
    heading.keep_with_next = True

    document.styles["Title"].paragraph_format.space_after = Pt(16)
    document.styles["Title"].paragraph_format.keep_with_next = True
    document.styles["Subtitle"].font.color.rgb = RGBColor(105, 112, 112)
    document.styles["Subtitle"].paragraph_format.keep_with_next = True
    document.add_paragraph("视频简介与文字稿", style="Title")
    if account:
        document.add_paragraph(f"视频号  {account}", style="Subtitle")
    layout = transcript_presentation(description, transcript)
    for label, paragraphs in (
        ("简介：", layout["description_paragraphs"]),
        ("文字稿：", layout["transcript_paragraphs"]),
    ):
        document.add_paragraph(label, style="Heading 1")
        for text in paragraphs:
            if text.strip():
                paragraph = document.add_paragraph(text.strip())
                paragraph.paragraph_format.first_line_indent = Pt(22)
                # Let Chinese paragraphs wrap naturally without full-word gaps.
                language = OxmlElement("w:lang")
                language.set(qn("w:eastAsia"), "zh-CN")
                paragraph.runs[0]._element.get_or_add_rPr().append(language)

    document.core_properties.title = "视频简介与文字稿"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def write_video_docx(output_path: Path, description: str, transcript: str, account: str = "") -> Path:
    """先在内存生成完整文档，再写入目标文件。"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(video_docx_bytes(description, transcript, account))
    return output_path
