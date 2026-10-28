"""真实 HTTP 验证历史文字稿预览、导出及沙盒排版迁移。"""

from io import BytesIO
import shutil
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
import httpx
import pytest
from sqlalchemy import delete, select, update

from platform_core.knowledge.video_transcriber.utils.transcript_layout import LAYOUT_VERSION, transcript_presentation
from platform_core.services import collection_artifact_service as artifacts
from platform_core.services import wxchannels_service as service
from platform_core.storage.postgres.manager import pg_manager
from platform_core.storage.postgres.models_business import Department, TaskRecord, User
from platform_core.utils.auth_utils import AuthUtils
from platform_core.workspace.filesystem import Workspace
from platform_core.workspace.paths import user_workspace_dir

pytestmark = [pytest.mark.asyncio, pytest.mark.integration]
BASE = "/api/knowledge/wxchannels"


async def test_historical_preview_download_archive_and_workspace():
    """旧记录不重转录，所有展示路径采用一致段落且维持用户隔离。"""
    identities = []
    task_id, oid = uuid4().hex, uuid4().hex
    title = "这段视频简介介绍完整背景，保留原始文字和标点。" * 28
    text = "这是已经完成识别的正文内容，其中包含 75% 和 3.14 等数字。" * 42
    account = "pytest 排版验收"
    legacy = Document()
    legacy.add_paragraph(title)
    legacy.add_paragraph(text)
    buffer = BytesIO()
    legacy.save(buffer)
    original = buffer.getvalue()
    storage = service.get_minio_client()
    async with pg_manager.get_async_session_context() as db:
        department = await db.scalar(select(Department).limit(1))
        for _ in range(2):
            uid = "pytest_layout_" + uuid4().hex
            user = User(uid=uid, username=uid, password_hash="no-login", role="admin", department_id=department.id)
            db.add(user)
            await db.flush()
            identities.append(
                (user.id, uid, {"Authorization": "Bearer " + AuthUtils.create_access_token({"sub": str(user.id)})})
            )
    uid, headers = identities[0][1:]
    try:
        await service.store_video_document(uid, oid, original, storage)
        async with pg_manager.get_async_session_context() as db:
            db.add(
                TaskRecord(
                    id=task_id,
                    name="pytest historical layout",
                    type=service.TASK_TYPE,
                    status="success",
                    payload={"operator_id": uid, "account": account},
                    result={
                        "items": [
                            {
                                "index": 0,
                                "oid": oid,
                                "title": title,
                                "text": text,
                                "status": "completed",
                                "filename": "sample.docx",
                                "object_key": service.video_object_key(uid, oid),
                            }
                        ]
                    },
                )
            )
        record = artifacts.publish_video(uid, oid, title, text, original, account)
        workspace = Workspace(uid)
        artifacts.write_artifact_file(
            workspace,
            record["import_path"],
            f"# {title}\n\n视频号：{account}\n\n视频 OID：{oid}\n\n## 文字稿\n\n{text}\n".encode(),
        )
        with artifacts.manifest(uid, "wxchannels") as state:
            state["items"][record["id"]].pop("layout_version")
            state["items"][record["id"]].pop("generated_hashes")
        expected = transcript_presentation(title, text)

        def check_word(content):
            """检查 HTTP 实际返回的 Word 文本和段落层级。"""
            paragraphs = [p.text for p in Document(BytesIO(content)).paragraphs]
            assert paragraphs == [
                "视频简介与文字稿",
                f"视频号  {account}",
                "简介：",
                *expected["description_paragraphs"],
                "文字稿：",
                *expected["transcript_paragraphs"],
            ]

        async with httpx.AsyncClient(base_url="http://127.0.0.1:5050", timeout=60) as client:
            response = await client.get(BASE + f"/tasks/{task_id}", headers=headers)
            assert response.status_code == 200, response.text
            item = response.json()["result"]["items"][0]
            assert "object_key" not in item
            for key, paragraphs in expected.items():
                assert len(paragraphs) > 2 and item[key] == paragraphs
            for suffix in ("", "/items/0/docx", "/archive"):
                url = BASE + f"/tasks/{task_id}" + suffix
                assert (await client.get(url, headers=identities[1][2])).status_code == 404
            response = await client.get(BASE + f"/tasks/{task_id}/items/0/docx", headers=headers)
            assert response.status_code == 200, response.text
            check_word(response.content)
            response = await client.get(BASE + f"/tasks/{task_id}/archive", headers=headers)
            assert response.status_code == 200, response.text
            check_word(ZipFile(BytesIO(response.content)).read("sample.docx"))
            response = await client.post("/api/knowledge/collection-artifacts/wxchannels/sync", headers=headers)
            assert response.status_code == 200, response.text
            assert not response.json()["errors"]
            record = response.json()["items"][0]
            assert record["layout_version"] == LAYOUT_VERSION
            response = await client.get(
                "/api/workspace/download", params={"path": record["primary_path"]}, headers=headers
            )
            assert response.status_code == 200, response.text
            check_word(response.content)
            response = await client.get(
                "/api/workspace/download", params={"path": record["import_path"]}, headers=headers
            )
            assert response.status_code == 200, response.text
            for paragraph in expected["transcript_paragraphs"]:
                assert "\n\n" + paragraph in response.text
            assert (
                await storage.adownload_file(storage.KB_BUCKETS["documents"], service.video_object_key(uid, oid))
                == original
            )
    finally:
        await storage.adelete_objects_by_prefix(storage.KB_BUCKETS["documents"], service.video_owner_prefix(uid))
        shutil.rmtree(artifacts.manifest_directory(uid), ignore_errors=True)
        shutil.rmtree(user_workspace_dir(uid).parent, ignore_errors=True)
        async with pg_manager.get_async_session_context() as db:
            await db.execute(delete(TaskRecord).where(TaskRecord.id == task_id))
            await db.execute(update(User).where(User.id.in_([i[0] for i in identities])).values(is_deleted=True))
