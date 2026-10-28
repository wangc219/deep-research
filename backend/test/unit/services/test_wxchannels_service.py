"""验证视频号转换的去重、故障隔离、取消和文件格式契约。"""

import asyncio
from io import BytesIO
from types import SimpleNamespace

import httpx
import pytest
from docx import Document
from fastapi import FastAPI

from platform_core.knowledge.video_transcriber.utils.docx_exporter import write_video_docx
from platform_core.knowledge.video_transcriber.utils.text_normalizer import strip_trailing_engagement_prompt
from platform_core.knowledge.wxchannels import WxChannelsError, WxChannelsMCPClient, WxVideo
from platform_core.services import wxchannels_service as service
from server.routers import wxchannels_router as api
from server.utils.auth_middleware import get_required_user

pytestmark = pytest.mark.asyncio


@pytest.fixture(autouse=True)
def isolated_collection_workspace(monkeypatch, tmp_path):
    """转换单测只写临时用户工作区和成果清单。"""
    monkeypatch.setenv("DEEP_RESEARCH_USER_DATA_DIR", str(tmp_path / "users"))
    monkeypatch.setenv("COLLECTION_ARTIFACT_DATA_ROOT", str(tmp_path / "artifacts"))


@pytest.fixture(autouse=True)
def memory_channel_index(monkeypatch):
    """用内存目录隔离单元测试，真实事务另由 PostgreSQL 集成测试覆盖。"""
    from unittest.mock import AsyncMock

    monkeypatch.setenv("WX_CHANNELS_SHARED_CLOUD", "false")
    repository = SimpleNamespace(find=AsyncMock(return_value=None))
    records = {}

    async def merge(owner, account, username, videos, *, complete):
        previous = len(records)
        records.update({v["oid"]: v for v in videos})
        return {
            "id": "channel",
            "username": username,
            "video_count": len(records),
            "new_count": len(records) - previous,
            "history_complete": complete,
        }

    async def videos(owner, channel, *, offset=0, limit=50):
        ordered = sorted(records.values(), key=lambda item: item["published_at"], reverse=True)
        return {"total_count": len(ordered), "videos": ordered[offset : offset + limit]}

    async def pending_videos(owner, channel, completed, limit):
        ordered = sorted(records.values(), key=lambda item: item["published_at"], reverse=True)
        pending = [v for v in ordered if v["oid"] not in completed]
        return {"total_count": len(ordered), "completed_count": len(ordered) - len(pending), "videos": pending[:limit]}

    repository.pending_videos = pending_videos
    repository.merge = merge
    repository.videos = videos
    monkeypatch.setattr(service, "WechatChannelRepository", lambda: repository)
    return repository


async def test_existing_collector_reuses_working_proxy_and_certificate(monkeypatch):
    """连接已工作的 video_text 时只检查，不重启代理或替换证书。"""
    requests = []

    def respond(request):
        requests.append(request)
        if request.url.path == "/api/proxy/status":
            return httpx.Response(
                200,
                json={
                    "code": 0,
                    "data": {
                        "certificate": {"installed": True, "trusted": True, "is_legacy": False},
                        "service": {"listening": True},
                        "system_proxy": {"matched": True},
                    },
                },
            )
        return httpx.Response(200, json={"code": 0, "data": {"enabled": True}})

    client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
    monkeypatch.setattr(service.httpx, "AsyncClient", lambda **kwargs: client)
    monkeypatch.setattr(service, "downloader_endpoint", lambda: "http://native:2022/mcp")

    async def status():
        return {"ready": True}

    monkeypatch.setattr(service, "connection_status", status)
    assert await service.connect_existing_collector() == {"ready": True}
    assert [(request.method, str(request.url)) for request in requests] == [
        ("POST", "http://native:2022/api/mcp/enable"),
        ("GET", "http://native:2022/api/proxy/status"),
    ]


class Context:
    """保存 checkpoint 快照，模拟平台 lease 提供的 Handler 契约。"""

    def __init__(self):
        self.task_id = "task"
        self.payload = {
            "operator_id": "alice",
            "account": "账号",
            "max_pages": 2,
            "max_videos": 2,
            "language": "zh",
            "timeout_seconds": 30,
        }
        self.result = None

    async def set_result(self, result):
        import copy

        self.result = copy.deepcopy(result)

    async def set_progress(self, *args):
        pass

    async def set_message(self, *args):
        pass

    async def raise_if_cancelled(self):
        pass


@pytest.mark.parametrize("oids, expected_timeout", [([], 21600), (["one-video"], 7800)])
async def test_default_http_submission_respects_worker_timeout(monkeypatch, oids, expected_timeout):
    """默认二十条与单选任务都可提交，不能因推算时限超过 Worker 上限报 500。"""
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[get_required_user] = lambda: SimpleNamespace(uid="alice")
    captured = []

    async def status(**kwargs):
        return {"ready": True}

    async def enqueue(**kwargs):
        api.tasker._resolve_timeout_seconds(kwargs["timeout_seconds"])
        captured.append(kwargs)
        return SimpleNamespace(id="task"), True

    monkeypatch.setattr(service, "connection_status", status)
    monkeypatch.setattr(api.tasker, "default_timeout_seconds", 21600)
    monkeypatch.setattr(api.tasker, "enqueue_unique_by_payload", enqueue)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post("/wxchannels/transcribe", json={"account": "火研所", "video_oids": oids})
    assert response.status_code == 202 and response.json()["submitted"]
    assert captured[0]["timeout_seconds"] == expected_timeout


class Storage:
    """记录私有文件内容，支持模拟已删除的文档。"""

    KB_BUCKETS = {"documents": "private"}

    def __init__(self):
        self.files = {}

    async def alist_object_metadata(self, bucket, prefix):
        return [{"object_name": key} for key in self.files if key.startswith(prefix)]

    async def astat_file(self, bucket, key):
        return len(self.files[key]) if key in self.files else None

    async def aupload_file(self, *, object_name, data, **kwargs):
        self.files[object_name] = data
        return SimpleNamespace(object_name=object_name)


@pytest.fixture
def pipeline(monkeypatch):
    """隔离外部下载器和识别模型，保留真实 DOCX 生成。"""
    storage = Storage()
    context = Context()
    history = []
    videos = [WxVideo(str(i), "nonce", title=f"简介{i}") for i in range(5)]
    extracted = []

    class Client:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            pass

        async def search_accounts(self, account):
            return [{"nickname": account, "username": "account@finder"}]

        async def iter_account_videos(self, *args, **kwargs):
            return videos

        async def extract_audio_stream(self, video, path, **kwargs):
            path.write_bytes(b"audio")
            extracted.append((video.oid, path))
            return path

    async def transcribe(audio, *args):
        return {"text": "正文需要关注工程约束。\n\n记得点赞、关注哦。", "language": "zh", "confidence": 0.9}

    async def owned(owner):
        return history

    monkeypatch.setattr(service, "WxChannelsMCPClient", lambda *a, **kw: Client())
    monkeypatch.setattr(service, "get_minio_client", lambda: storage)
    monkeypatch.setattr(service, "transcribe_audio", transcribe)
    monkeypatch.setattr(service, "list_owned_tasks", owned)
    return context, storage, history, extracted


async def test_filters_completed_before_applying_new_video_limit(pipeline):
    context, storage, history, extracted = pipeline
    key = service.video_object_key("alice", "0")
    storage.files[key] = b"existing"
    storage.files[service.video_completion_key("alice", "0")] = b"{}"
    history.append({"id": "previous", "result": {"items": [{"oid": "0", "status": "completed"}]}})
    result = await service.run_wxchannels_transcribe(context)
    assert [oid for oid, _ in extracted] == ["1", "2"]
    assert result["completed_count"] == 2
    assert result["skipped_count"] == 1
    assert not any(path.exists() for _, path in extracted)
    word = Document(BytesIO(storage.files[service.video_object_key("alice", "1")]))
    text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert "简介：" in text and "文字稿：" in text and "简介1" in text
    assert "正文需要关注工程约束。" in text and "记得点赞" not in text


async def test_orphaned_object_without_checkpoint_is_reprocessed(pipeline):
    context, storage, _, extracted = pipeline
    storage.files[service.video_object_key("alice", "0")] = b"orphan"
    await service.run_wxchannels_transcribe(context)
    assert extracted[0][0] == "0"


async def test_deleted_docx_is_eligible_for_retry(pipeline):
    context, _, history, extracted = pipeline
    history.append({"id": "old", "result": {"items": [{"oid": "0", "status": "completed"}]}})
    await service.run_wxchannels_transcribe(context)
    assert extracted[0][0] == "0"


async def test_one_failure_does_not_stop_other_videos(pipeline, monkeypatch):
    context, storage, _, extracted = pipeline
    calls = 0

    async def transcribe(*args):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise RuntimeError("temporary failure")
        return {"text": "有效正文", "language": "zh", "confidence": 0.8}

    monkeypatch.setattr(service, "transcribe_audio", transcribe)
    result = await service.run_wxchannels_transcribe(context)
    assert result["failed_count"] == result["completed_count"] == 1
    assert len(storage.files) == 2
    assert not any(path.exists() for _, path in extracted)


async def test_all_failures_fail_the_platform_task(pipeline, monkeypatch):
    context, _, _, extracted = pipeline

    async def fail(*args):
        raise RuntimeError("asr unavailable")

    monkeypatch.setattr(service, "transcribe_audio", fail)
    with pytest.raises(WxChannelsError, match="均处理失败"):
        await service.run_wxchannels_transcribe(context)
    assert context.result["failed_count"] == 2
    assert not any(path.exists() for _, path in extracted)


async def test_cancel_cleans_audio_and_records_interrupted_item(pipeline, monkeypatch):
    context, storage, _, extracted = pipeline

    async def cancel(*args):
        raise asyncio.CancelledError()

    monkeypatch.setattr(service, "transcribe_audio", cancel)
    with pytest.raises(asyncio.CancelledError):
        await service.run_wxchannels_transcribe(context)
    assert context.result["items"][0]["status"] == "cancelled"
    assert not storage.files
    assert not any(path.exists() for _, path in extracted)


async def test_selected_videos_must_still_exist(pipeline):
    context, storage, _, _ = pipeline
    context.payload["video_oids"] = ["missing"]
    with pytest.raises(WxChannelsError, match="重新获取"):
        await service.run_wxchannels_transcribe(context)
    assert not storage.files


async def test_account_name_ambiguity_requires_finder_id():
    client = SimpleNamespace()

    async def search(account):
        return [{"nickname": account, "username": "a@finder"}, {"nickname": account, "username": "b@finder"}]

    client.search_accounts = search
    with pytest.raises(WxChannelsError, match="多个同名"):
        await service.resolve_username(client, "同名")
    assert await service.resolve_username(client, "stable@finder") == "stable@finder"


async def test_pagination_deduplicates_oid_despite_changed_nonce():
    client = WxChannelsMCPClient()
    calls = []

    async def page(username, marker):
        calls.append(marker)
        return {"objects": [{"id": "same", "objectNonceId": str(len(calls))}], "next_marker": "repeat"}

    client.account_videos = page
    result = await client.iter_account_videos("account@finder", 10)
    assert len(result) == 1 and len(calls) == 2


async def test_asr_cancel_kills_model_process(monkeypatch, tmp_path):
    original = asyncio.create_subprocess_exec
    started = asyncio.Event()
    processes = []

    async def spawn(*args, **kwargs):
        import sys

        process = await original(sys.executable, "-c", "import time; time.sleep(60)", start_new_session=True)
        processes.append(process)
        started.set()
        return process

    monkeypatch.setattr(asyncio, "create_subprocess_exec", spawn)
    future = asyncio.create_task(service.transcribe_audio(tmp_path / "a.wav", tmp_path / "out.json", "zh", 60))
    await started.wait()
    future.cancel()
    with pytest.raises(asyncio.CancelledError):
        await future
    assert processes[0].returncode is not None


@pytest.mark.parametrize("endpoint", ["/tasks/stolen", "/tasks/stolen/items/0/docx", "/tasks/stolen/archive"])
async def test_http_denies_cross_user_state_and_download(monkeypatch, endpoint):
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[get_required_user] = lambda: SimpleNamespace(uid="bob")

    async def get_task(task_id):
        return {"type": service.TASK_TYPE, "payload": {"operator_id": "alice"}}

    monkeypatch.setattr(api.tasker, "get_task", get_task)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        response = await client.get("/wxchannels" + endpoint)
    assert response.status_code == 404


async def test_http_denies_cross_user_cancel(monkeypatch):
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[get_required_user] = lambda: SimpleNamespace(uid="bob")

    async def get_task(task_id):
        return {"type": service.TASK_TYPE, "payload": {"operator_id": "alice"}}

    monkeypatch.setattr(api.tasker, "get_task", get_task)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post("/wxchannels/tasks/stolen/cancel")
    assert response.status_code == 404


@pytest.mark.parametrize("text", ["正文需要关注变化。", "关注系统约束是研究重点。"])
async def test_cleaner_preserves_normal_use_of_follow_word(text):
    assert strip_trailing_engagement_prompt(text) == text


async def test_document_contains_real_editable_paragraphs(tmp_path):
    path = write_video_docx(tmp_path / "out.docx", "原简介", "第一段。\n\n第二段。")
    paragraphs = [paragraph.text for paragraph in Document(path).paragraphs]
    assert paragraphs[-2:] == ["第一段。", "第二段。"]
    assert service.video_object_key("alice", "id") != service.video_object_key("bob", "id")
    assert "/" not in service.video_filename("../标题", "id")


async def test_all_pages_passes_old_limit_and_sorts_publication_time():
    """默认遍历超过一百页，置顶旧视频和毫秒时间戳不改变最新排序。"""
    client = WxChannelsMCPClient()
    calls = []

    async def page(username, marker):
        index = int(marker or 0)
        calls.append(index)
        return {
            "objects": [
                {"id": str(index), "objectNonceId": "n", "createtime": 1700000000 + index},
                {"id": "pinned", "objectNonceId": str(index), "createTime": 1600000000000},
                {"id": "unknown", "objectNonceId": "n", "createtime": "invalid"},
            ],
            "next_marker": str(index + 1) if index < 102 else "",
        }

    client.account_videos = page
    result = await client.iter_account_videos("account@finder")
    assert len(calls) == 103
    assert len(result) == 105
    assert result[0].oid == "102"
    assert result[-2].oid == "pinned" and result[-2].published_at == 1600000000
    assert result[-1].oid == "unknown" and result[-1].published_at == 0
    calls.clear()
    result = await client.iter_account_videos("account@finder", 2)
    assert len(calls) == 2 and result[0].oid == "1"


async def test_latest_pending_videos_ignore_pins_and_old_task_history(pipeline, monkeypatch):
    """先按时间排序和永久完成标记去重，再取最新 N 条，重复任务不会再处理同一视频。"""
    context, storage, _, extracted = pipeline
    original = service.WxChannelsMCPClient

    def client_factory(*args, **kwargs):
        client = original(*args, **kwargs)

        async def videos(*args, **kwargs):
            return [
                WxVideo(str(i), "n", title=f"视频{i}", published_at=t) for i, t in enumerate([100, 500, 300, 400, 200])
            ]

        client.iter_account_videos = videos
        return client

    monkeypatch.setattr(service, "WxChannelsMCPClient", client_factory)
    context.payload["max_pages"] = None
    await service.store_video_document("alice", "1", b"already completed", storage)
    result = await service.run_wxchannels_transcribe(context)
    assert [oid for oid, _ in extracted] == ["3", "2"]
    assert result["skipped_count"] == 1
    discovered = await service.discover_videos("alice", "账号")
    assert [item["oid"] for item in discovered["videos"]] == ["4", "0"]
    assert discovered["completed_count"] == 3
    assert discovered["videos"][0]["published_at"] == 200
    await service.run_wxchannels_transcribe(context)
    assert [oid for oid, _ in extracted] == ["3", "2", "4", "0"]
    result = await service.run_wxchannels_transcribe(context)
    assert result["completed_count"] == 0 and result["skipped_count"] == 5


async def test_incremental_feed_skips_two_known_pages_but_not_a_pinned_page():
    """已知置顶页之后仍发现新视频，连续两页已知历史才停止。"""
    client = WxChannelsMCPClient()
    calls = []
    pages = [["pin"], ["new", "old"], ["old"], ["older"], ["unneeded"]]

    async def page(username, marker):
        index = int(marker or 0)
        calls.append(index)
        return {
            "objects": [{"id": oid, "objectNonceId": "n", "createtime": 100 - index} for oid in pages[index]],
            "next_marker": str(index + 1),
        }

    client.account_videos = page
    result = await client.iter_account_videos("test@finder", known_oids={"pin", "old", "older"})
    assert calls == [0, 1, 2, 3]
    assert {video.oid for video in result} == {"pin", "new", "old", "older"}
    assert client.history_complete is False


@pytest.mark.parametrize("count", [5, 10, 15, 20, None])
async def test_discovery_latest_range_uses_cached_history_after_completed_filter(pipeline, monkeypatch, count):
    """5 条范围先跳过已完成内容，全部不会隐含截断。"""
    _, storage, _, _ = pipeline
    original = service.WxChannelsMCPClient

    def factory(*args, **kwargs):
        client = original(*args, **kwargs)

        async def videos(*args, **kwargs):
            return [WxVideo(str(i), "n", title=f"视频{i}", published_at=i) for i in range(30)]

        client.iter_account_videos = videos
        return client

    monkeypatch.setattr(service, "WxChannelsMCPClient", factory)
    for oid in ("29", "28"):
        await service.store_video_document("alice", oid, b"done", storage)
    result = await service.discover_videos("alice", "账号", max_videos=count)
    assert [item["oid"] for item in result["videos"]] == [str(i) for i in range(27, -1, -1)][:count]
    assert result["completed_count"] == 2


async def test_all_transcription_processes_more_than_one_hundred(pipeline, monkeypatch):
    """全部的真实流水线不被旧的 100 条设置上限静默截断。"""
    context, _, _, extracted = pipeline
    original = service.WxChannelsMCPClient

    def factory(*args, **kwargs):
        client = original(*args, **kwargs)

        async def videos(*args, **kwargs):
            return [WxVideo(str(i), "n", title=f"视频{i}", published_at=i) for i in range(101)]

        client.iter_account_videos = videos
        return client

    monkeypatch.setattr(service, "WxChannelsMCPClient", factory)
    context.payload["max_videos"] = None
    result = await service.run_wxchannels_transcribe(context)
    assert result["completed_count"] == len(extracted) == 101
    assert extracted[0][0] == "100" and extracted[-1][0] == "0"


async def test_managed_cloud_discovery_is_owned_and_queued(monkeypatch):
    """普通用户请求进入共享采集队列，不能伪造其他用户身份。"""
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[get_required_user] = lambda: SimpleNamespace(uid="alice")
    monkeypatch.setenv("WX_CHANNELS_SHARED_CLOUD", "true")
    calls = []

    async def enqueue(**kwargs):
        calls.append(kwargs)
        return SimpleNamespace(id="central-task"), True

    monkeypatch.setattr(api.tasker, "enqueue_unique_by_payload", enqueue)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post("/wxchannels/videos", json={"account": "test@finder", "operator_id": "mallory"})
    assert response.status_code == 200 and response.json()["task_id"] == "central-task"
    assert calls[0]["task_type"] == service.CENTRAL_DISCOVERY_TASK_TYPE
    assert calls[0]["payload"]["operator_id"] == "alice"
    assert calls[0]["payload_match"]["operator_id"] == "alice"
    assert calls[0]["payload_match"]["account"] == "test@finder"


async def test_shared_cloud_discovery_handler_preserves_user_range(monkeypatch):
    """持久任务恢复后仍使用请求用户与范围，而非共享微信身份保存目录。"""
    from platform_core.services.task_registry import get_task_definition

    context = Context()
    context.payload.update(max_videos=15, full_refresh=True)
    calls = []

    async def discover(*args):
        calls.append(args)
        return {"videos": [], "new_count": 0}

    monkeypatch.setattr(service, "discover_videos", discover)
    handler = get_task_definition(service.CENTRAL_DISCOVERY_TASK_TYPE).load_handler()
    result = await handler(context)
    assert result == {"videos": [], "new_count": 0}
    assert calls == [("alice", "账号", 2, 15, True, False, None)]


@pytest.mark.parametrize(
    "fetched_at,full_refresh,upstream", [(200, False, False), (50, False, True), (200, True, True)]
)
async def test_shared_page_snapshot_preserves_refresh_and_owner(monkeypatch, fetched_at, full_refresh, upstream):
    """同批用户共享公开页，新的显式刷新与完整同步始终访问源。"""
    import json
    from unittest.mock import AsyncMock
    from platform_core.services import run_queue_service

    snapshot = {
        "fetched_at": fetched_at,
        "history_complete": False,
        "username": "public@finder",
        "videos": [{"oid": "public", "nid": "n", "title": "公开视频", "published_at": 1}],
        "next_marker": "page2",
    }
    pool = SimpleNamespace(
        get=AsyncMock(return_value=json.dumps(snapshot)),
        set=AsyncMock(),
        hset=AsyncMock(),
        expire=AsyncMock(),
        exists=AsyncMock(return_value=False),
    )
    monkeypatch.setattr(run_queue_service, "get_arq_pool", AsyncMock(return_value=pool))
    monkeypatch.setenv("WX_CHANNELS_SHARED_CLOUD", "true")
    repo = SimpleNamespace(find=AsyncMock(return_value=None), merge=AsyncMock(return_value={"id": "bob-channel"}))
    monkeypatch.setattr(service, "WechatChannelRepository", lambda: repo)
    own = AsyncMock(return_value={"owner": "bob"})
    monkeypatch.setattr(service, "saved_video_discovery", own)
    calls = []

    class Client:
        history_complete = False
        next_marker = "page2"

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            pass

        async def search_accounts(self, account):
            return [{"nickname": account, "username": "public@finder"}]

        async def iter_account_videos(self, username, pages):
            calls.append(pages)
            assert repo.merge.called, "账号必须在分页抓取之前保存"
            return [WxVideo("public", "n", title="公开视频", published_at=1)]

    monkeypatch.setattr(service, "WxChannelsMCPClient", lambda *a, **k: Client())
    assert await service.discover_videos("bob", "账号", full_refresh=full_refresh, requested_at=100) == {"owner": "bob"}
    assert len(calls) == int(upstream)
    assert repo.merge.await_args.args[0] == "bob"
    if upstream:
        assert calls == [None if full_refresh else 2]
    if upstream and not full_refresh:
        saved = json.loads(pool.set.await_args_list[0].args[1])
        assert set(saved) == {"fetched_at", "history_complete", "username", "videos", "next_marker"}
        assert "operator_id" not in saved


async def test_more_history_continues_owned_cursor(monkeypatch):
    """加载更多从本人保存的游标继续，不从首页重扫；刷新最新不覆盖历史游标。"""
    from unittest.mock import AsyncMock
    from platform_core.services import run_queue_service

    pool = SimpleNamespace(
        get=AsyncMock(side_effect=[b'{"marker":"next-private-page","pages":2}', None]),
        set=AsyncMock(),
        hset=AsyncMock(),
        expire=AsyncMock(),
        exists=AsyncMock(return_value=True),
    )
    monkeypatch.setattr(run_queue_service, "get_arq_pool", AsyncMock(return_value=pool))
    monkeypatch.setenv("WX_CHANNELS_SHARED_CLOUD", "true")
    repo = SimpleNamespace(
        find=AsyncMock(return_value={"id": "alice-channel", "username": "public@finder", "history_complete": False}),
        merge=AsyncMock(return_value={"id": "alice-channel"}),
    )
    monkeypatch.setattr(service, "WechatChannelRepository", lambda: repo)
    monkeypatch.setattr(service, "saved_video_discovery", AsyncMock(return_value={"new_count": 1}))

    class Client:
        history_complete = False
        next_marker = "later-page"

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            pass

        async def iter_account_videos(self, username, pages, **kwargs):
            assert pages == 2 and kwargs == {"start_marker": "next-private-page"}
            return [WxVideo("new-video", "nonce", published_at=1)]

    monkeypatch.setattr(service, "WxChannelsMCPClient", lambda *a, **k: Client())
    await service.discover_videos("alice", "账号", load_more=True)
    assert pool.get.await_args_list[0].args == ("wxchannels:cursor:v1:alice-channel",)
    import json

    assert pool.set.await_args_list[-1].args[0] == "wxchannels:cursor:v1:alice-channel"
    assert json.loads(pool.set.await_args_list[-1].args[1]) == {"marker": "later-page", "pages": 4}


async def test_feed_stops_at_requested_pending_count_and_selected_video():
    """即使账号有十万页，最新150条及指定旧视频找到后立即停止。"""
    client = WxChannelsMCPClient()
    calls = []

    async def page(username, marker):
        index = int(marker or 0)
        calls.append(index)
        return {
            "objects": [
                {"id": str(index * 15 + i), "objectNonceId": "n", "createtime": 100000 - index * 15 - i}
                for i in range(15)
            ],
            "next_marker": str(index + 1),
        }

    client.account_videos = page
    videos = await client.iter_account_videos("account@finder", max_videos=150)
    assert len(videos) == 150 and len(calls) == 10 and client.next_marker == "10"
    calls.clear()
    videos = await client.iter_account_videos("account@finder", selected_oids={"31"})
    assert len(calls) == 3 and any(v.oid == "31" for v in videos)
    calls.clear()
    videos = await client.iter_account_videos("account@finder", max_videos=5, skip_oids={str(i) for i in range(15)})
    assert len(calls) == 2


@pytest.mark.parametrize("count", [30, 40, 50, 100, 150])
async def test_extended_range_is_accepted_and_bounded(count):
    """新增数量选项由后端模型接纳，不被旧100条上限拒绝。"""
    assert api.VideosRequest(account="账号", max_videos=count).max_videos == count
    assert api.TranscribeRequest(account="账号", max_videos=count).max_videos == count
