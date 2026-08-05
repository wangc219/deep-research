from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook


QUERY_HEADERS = {
    "query",
    "研究方向",
    "研究问题",
    "问题",
    "选题",
    "主题",
    "需求query",
}
SUPPLEMENT_HEADERS = {
    "supplemental_information",
    "补充信息",
    "补充材料",
    "核心研究重点/制胜机理",
    "核心研究重点",
    "制胜机理",
    "研究重点",
    "问题描述",
}
CATEGORY_HEADERS = {"一级领域", "分类", "类别", "领域", "方向分类"}
PRIORITY_HEADERS = {"重点层级", "优先级", "层级", "重要程度"}
RATIONALE_HEADERS = {"generation_rationale", "生成理由", "研究理由", "立项理由"}
NUMBERED_ITEM = re.compile(r"(?m)^\s*(\d{1,3})[、.．]\s*(?=\S)")
SECTION_HEADING = re.compile(r"^\s*[一二三四五六七八九十百]+、[^\n]+\s*$", re.MULTILINE)


def parse_query_file(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        return _parse_xlsx(filename, content)
    if suffix == ".csv":
        return _parse_csv(filename, content)
    raise ValueError("only .xlsx and .csv Query lists are supported")


def parse_numbered_query_text(
    content: str, *, source_name: str
) -> list[dict[str, Any]]:
    cleaned = SECTION_HEADING.sub("\n", content.replace("\r\n", "\n").strip())
    matches = list(NUMBERED_ITEM.finditer(cleaned))
    raw_items: list[str]
    if matches:
        raw_items = []
        for index, match in enumerate(matches):
            end = (
                matches[index + 1].start() if index + 1 < len(matches) else len(cleaned)
            )
            raw_items.append(cleaned[match.end() : end].strip())
    else:
        raw_items = [
            item.strip() for item in re.split(r"\n\s*\n+", cleaned) if item.strip()
        ]
        if len(raw_items) == 1:
            raw_items = [item.strip() for item in cleaned.splitlines() if item.strip()]

    records: list[dict[str, Any]] = []
    for index, item in enumerate(raw_items, start=1):
        query = re.sub(r"\s+", " ", item).strip()
        if not query:
            continue
        category_match = re.match(r"^【([^】]+)】", query)
        category = category_match.group(1).strip() if category_match else ""
        rationale = "人工批量导入的长 Query"
        if category:
            rationale += f"；分类：{category}"
        records.append(
            {
                "query": query,
                "supplemental_information": "",
                "generation_rationale": rationale,
                "source_references": [
                    _document_reference(
                        source_name,
                        f"长 Query 文本第 {index} 项",
                    )
                ],
                "source_row": index,
            }
        )
    return records


def _parse_xlsx(filename: str, content: bytes) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"unable to read xlsx workbook: {exc}") from exc
    records: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        if not rows:
            continue
        header_index, headers = _find_header(rows)
        if header_index is None:
            continue
        for row_number, row in enumerate(
            rows[header_index + 1 :], start=header_index + 2
        ):
            record = _row_to_record(
                headers,
                row,
                filename=filename,
                location=f"工作表“{worksheet.title}”第 {row_number} 行",
                source_row=row_number,
            )
            if record:
                records.append(record)
    if not records:
        raise ValueError("no Query rows found; include a ‘研究方向’ or ‘Query’ column")
    return records


def _parse_csv(filename: str, content: bytes) -> list[dict[str, Any]]:
    decoded = _decode_csv(content)
    rows = list(csv.reader(StringIO(decoded)))
    header_index, headers = _find_header(rows)
    if header_index is None:
        raise ValueError("no Query column found in csv")
    records = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        record = _row_to_record(
            headers,
            row,
            filename=filename,
            location=f"CSV 第 {row_number} 行",
            source_row=row_number,
        )
        if record:
            records.append(record)
    if not records:
        raise ValueError("no non-empty Query rows found in csv")
    return records


def _find_header(rows: list[list[Any]]) -> tuple[int | None, list[str]]:
    for index, row in enumerate(rows[:30]):
        headers = [_normalize_header(value) for value in row]
        if any(item in QUERY_HEADERS for item in headers):
            return index, headers
    return None, []


def _row_to_record(
    headers: list[str],
    row: list[Any],
    *,
    filename: str,
    location: str,
    source_row: int,
) -> dict[str, Any] | None:
    values = {
        header: _cell_text(row[index]) if index < len(row) else ""
        for index, header in enumerate(headers)
        if header
    }
    query = _first(values, QUERY_HEADERS)
    if not query:
        return None
    supplement = _first(values, SUPPLEMENT_HEADERS)
    rationale_parts = []
    explicit_rationale = _first(values, RATIONALE_HEADERS)
    category = _first(values, CATEGORY_HEADERS)
    priority = _first(values, PRIORITY_HEADERS)
    if explicit_rationale:
        rationale_parts.append(explicit_rationale)
    if category:
        rationale_parts.append(f"一级领域：{category}")
    if priority:
        rationale_parts.append(f"重点层级：{priority}")
    return {
        "query": query,
        "supplemental_information": supplement,
        "generation_rationale": "；".join(rationale_parts) or "人工文件导入 Query",
        "source_references": [_document_reference(filename, location)],
        "source_row": source_row,
    }


def _document_reference(title: str, relevance_note: str) -> dict[str, str]:
    return {
        "title": title,
        "url": "",
        "relevance_note": relevance_note,
        "source_kind": "document",
    }


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", _cell_text(value)).lower()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _first(values: dict[str, str], candidates: set[str]) -> str:
    for header, value in values.items():
        if header in candidates and value:
            return value
    return ""


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("csv must use UTF-8 or GB18030 encoding")
