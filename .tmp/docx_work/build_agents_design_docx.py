from __future__ import annotations

from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/wangchen/equipment research")
SOURCE = ROOT / "docs/architecture/二级业务智能体_国际形势与作战场景_Prompt_Tools_Skills_Harness设计.md"
OUTPUT = ROOT / "docs/architecture/二级业务智能体_国际形势与作战场景_Prompt_Tools_Skills_Harness设计.docx"
ASSET_DIR = ROOT / ".tmp/docx_work/assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)
FLOW_IMAGE = ASSET_DIR / "agent-collaboration-flow.png"


FONT_NAME = "Hiragino Sans GB"
MONO_NAME = "Hiragino Sans GB"
NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEXT = "20242A"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CODE_FILL = "F6F8FA"
BORDER = "C9D2DE"
WHITE = "FFFFFF"


def set_run_font(run, *, name=FONT_NAME, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_paragraph_border_bottom(paragraph, color=BORDER, size="6", space="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_number(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.1

    for name, size, color, before, after in (
        ("Title", 24, NAVY, 0, 6),
        ("Subtitle", 12, MUTED, 0, 16),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, DARK_BLUE, 9, 4),
    ):
        style = doc.styles[name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = name.startswith("Heading")
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.1


def draw_flow_image(path):
    width, height = 1800, 470
    image = Image.new("RGB", (width, height), f"#{WHITE}")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Hiragino Sans GB.ttc"
    title_font = ImageFont.truetype(font_path, 33)
    body_font = ImageFont.truetype(font_path, 26)
    small_font = ImageFont.truetype(font_path, 21)
    boxes = [
        (50, 130, 330, 320, "Chief Orchestrator", "任务语义卡与边界"),
        (395, 100, 700, 350, "国际形势分析 Agent", "3—5个背景假设"),
        (765, 100, 1070, 350, "StrategicAssessment", "证据/触发器/反证"),
        (1135, 100, 1440, 350, "作战场景推演 Agent", "每背景2—3个场景"),
        (1505, 130, 1750, 320, "下游 Agent", "装备/运用/制胜"),
    ]
    for idx, (x1, y1, x2, y2, title, subtitle) in enumerate(boxes):
        fill = LIGHT_BLUE if idx in (1, 3) else LIGHT_GRAY
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=f"#{fill}", outline=f"#{BLUE}", width=4)
        bbox = draw.textbbox((0, 0), title, font=title_font if idx in (1, 3) else body_font)
        tx = x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2
        draw.text((tx, y1 + 55), title, font=title_font if idx in (1, 3) else body_font, fill=f"#{NAVY}")
        bbox2 = draw.textbbox((0, 0), subtitle, font=small_font)
        tx2 = x1 + (x2 - x1 - (bbox2[2] - bbox2[0])) / 2
        draw.text((tx2, y1 + 125), subtitle, font=small_font, fill=f"#{MUTED}")
        if idx < len(boxes) - 1:
            nx1 = boxes[idx + 1][0]
            y = 225
            draw.line((x2 + 12, y, nx1 - 22, y), fill=f"#{BLUE}", width=5)
            draw.polygon([(nx1 - 22, y - 12), (nx1 - 22, y + 12), (nx1 - 5, y)], fill=f"#{BLUE}")
    draw.text((40, 410), "Typed handoff · SavePoint · Recall from earliest failed node", font=small_font, fill=f"#{MUTED}")
    image.save(path)


def parse_inline(paragraph, text, *, size=10.5, color=TEXT):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name=MONO_NAME, size=size - 0.5, color=DARK_BLUE)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, CODE_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("\n".join(lines))
    set_run_font(run, name=MONO_NAME, size=8.3, color="24292F")
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def widths_for_table(headers, rows):
    n = len(headers)
    lowered = [h.lower() for h in headers]
    if n == 2:
        if headers[0] in {"项目", "设计项"}:
            return [2000, 7360]
        return [3300, 6060]
    if n == 5:
        return [1300, 1900, 2300, 1800, 2060]
    if n == 4:
        return [1600, 2700, 2500, 2560]
    if n == 3:
        return [1800, 3600, 3960]
    base = 9360 // n
    widths = [base] * n
    widths[-1] += 9360 - sum(widths)
    return widths


def add_markdown_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    widths = widths_for_table(headers, rows)
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        parse_inline(p, text, size=8.8, color=NAVY)
        for run in p.runs:
            run.bold = True
    for r_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            if r_index % 2:
                shade_cell(cells[idx], "FAFBFC")
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            parse_inline(p, text, size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_header_footer(doc):
    for section in doc.sections:
        section.header_distance = Inches(0.45)
        section.footer_distance = Inches(0.45)
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = hp.add_run("JS装备需求深度挖掘 · 二级业务智能体专用设计")
        set_run_font(run, size=8.5, color=MUTED)
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_page_number(fp)


def add_title_block(doc):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(5)
    run = kicker.add_run("TECHNICAL DESIGN · CODEX CLI AGENT RUNTIME")
    set_run_font(run, size=9.5, bold=True, color=BLUE)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(5)
    parse_inline(title, "二级业务智能体专用设计", size=24, color=NAVY)

    subtitle = doc.add_paragraph(style="Subtitle")
    parse_inline(subtitle, "国际形势分析 Agent 与作战场景推演 Agent\nPrompt · Tools · Skills · Harness 编排", size=12.5, color=MUTED)

    metadata = [
        ("版本", "v1.0"),
        ("日期", "2026-07-17"),
        ("执行器", "Codex CLI（隔离、ephemeral、Skill 驱动）"),
        ("依据", "架构设计 DOCX 与现有 equipment_deep_research 运行时"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    set_table_geometry(table, [1500, 7860])
    for idx, (label, value) in enumerate(metadata):
        left, right = table.rows[idx].cells
        shade_cell(left, LIGHT_BLUE)
        lp = left.paragraphs[0]
        parse_inline(lp, label, size=9, color=NAVY)
        lp.runs[0].bold = True
        rp = right.paragraphs[0]
        parse_inline(rp, value, size=9.2)
    spacer = doc.add_paragraph()
    set_paragraph_border_bottom(spacer, color=BLUE, size="10", space="6")
    spacer.paragraph_format.space_after = Pt(12)


def build_document():
    draw_flow_image(FLOW_IMAGE)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    configure_styles(doc)
    add_header_footer(doc)
    add_title_block(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 5  # skip source title metadata already represented in the opening block
    in_code = False
    code_lang = ""
    code_lines = []
    flow_inserted = False

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lang = stripped[3:].strip()
                code_lines = []
            else:
                if code_lang == "mermaid" and not flow_inserted:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    picture = p.add_run().add_picture(str(FLOW_IMAGE), width=Inches(6.45))
                    picture._inline.docPr.set(
                        "descr",
                        "Chief Orchestrator到国际形势分析Agent、StrategicAssessment、作战场景推演Agent和下游Agent的typed handoff流程",
                    )
                    cap = doc.add_paragraph("图 1  背景假设到作战场景的 typed handoff 主链")
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.paragraph_format.space_after = Pt(8)
                    for run in cap.runs:
                        set_run_font(run, size=8.5, color=MUTED, italic=True)
                    flow_inserted = True
                else:
                    add_code_block(doc, code_lines)
                in_code = False
                code_lang = ""
                code_lines = []
            index += 1
            continue

        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[index + 1].strip()):
            headers = [c.strip() for c in stripped.strip("|").split("|")]
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([c.strip() for c in lines[index].strip().strip("|").split("|")])
                index += 1
            add_markdown_table(doc, headers, rows)
            continue

        if not stripped:
            index += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            display_level = min(level, 3)
            text = heading.group(2)
            if text.startswith("二级业务智能体专用设计") or text.startswith("国际形势分析 Agent 与作战场景推演 Agent"):
                index += 1
                continue
            p = doc.add_paragraph(style=f"Heading {display_level}")
            parse_inline(
                p,
                text,
                size={1: 16, 2: 13, 3: 11.5}[display_level],
                color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[display_level],
            )
            index += 1
            continue

        bullet = re.match(r"^-\s+(.*)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            parse_inline(p, bullet.group(1))
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            parse_inline(p, numbered.group(1))
            index += 1
            continue

        p = doc.add_paragraph()
        parse_inline(p, stripped)
        index += 1

    core = doc.core_properties
    core.title = "二级业务智能体专用设计：国际形势分析与作战场景推演"
    core.subject = "Codex CLI Prompt、Tools、Skills 与 Harness 编排"
    core.author = "Codex"
    core.keywords = "Codex CLI, Agent, Harness, Prompt, Skill, Tool, JS装备需求"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
